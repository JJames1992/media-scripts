#!/usr/bin/env python3
"""Canonicalize the 'To See' sheet against TMDB.

Every row gets a deterministic identity for seed hunting:
canonical title, canonical year, media type (movie/series),
TMDB id, season count (series), and an aria2-ready search core.

Resolution order:
  1. IMDb id (already in the sheet) -> TMDB /find/{imdb_id}
  2. Fallback: TMDB /search/movie + /search/tv by title+year

If the TMDB title diverges from the sheet title (sim < 0.55),
the row is flagged `imdb_mismatch` instead of silently trusted —
that's exactly the "what is what" hiccup this kills.

Writes canonical columns back into Movies.xlsx (cols F..N) and a
machine-readable manifest to media-scripts/media_catalog.json.
"""
import difflib
import json
import re
import time
import urllib.parse
import urllib.request

import openpyxl

TMDB_KEY = "72e25a0a54b6c3ffbc9364bfc3e67287"
XLSX = "C:/Users/jonat/OneDrive/Movies.xlsx"
MANIFEST = "C:/Users/jonat/media-scripts/media_catalog.json"

NEW_COLS = [
    "Type", "Canonical Title", "Canonical Year", "TMDB ID",
    "IMDb ID", "Seasons", "Search Core", "Resolve Status", "Resolve Note",
]


def api(path, **params):
    params["api_key"] = TMDB_KEY
    url = f"https://api.themoviedb.org/3{path}?{urllib.parse.urlencode(params)}"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=20) as r:
        return json.load(r)


def imdb_id_from(link):
    if not link:
        return None
    m = re.search(r"tt\d{5,10}", str(link))
    return m.group(0) if m else None


def clean_title(t):
    t = re.sub(r"\s*\(TV [Ss]how\)\s*$", "", t.strip())
    t = re.sub(r"\s*\(\d{4}\)\s*$", "", t)
    return t.strip()


def similar(a, b):
    return difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()


def build_rec(tmdb, mtype, imdb):
    title = tmdb.get("title") or tmdb.get("name")
    if mtype == "movie":
        year = (tmdb.get("release_date") or "")[:4]
    else:
        year = (tmdb.get("first_air_date") or "")[:4]
    year = int(year) if year.isdigit() else ""
    return {
        "title": title,
        "year": year,
        "type": mtype,
        "tmdb_id": tmdb["id"],
        "imdb_id": imdb or "",
        "search_core": f"{title} ({year})" if year else title,
    }


def resolve(title, year, imdb):
    """-> (record, status, note). status: imdb|imdb_mismatch|search|failed"""
    sheet_clean = clean_title(title)
    # 1) IMDb id
    if imdb:
        try:
            d = api(f"/find/{imdb}", external_source="imdb_id")
            pool = d.get("movie_results") or d.get("tv_results") or []
            if pool:
                best = pool[0]
                mtype = "movie" if best.get("media_type") == "movie" else "series"
                rec = build_rec(best, mtype, imdb)
                sim = similar(sheet_clean, rec["title"])
                if sim < 0.55:
                    return rec, "imdb_mismatch", f"sheet title vs TMDB title sim={sim:.2f}"
                return rec, "imdb", ""
        except Exception:
            pass
    # 2) search fallback
    try:
        yr = int(float(year)) if year else None
        sm = api("/search/movie", query=sheet_clean, year=yr, include_adult="false") if yr \
            else api("/search/movie", query=sheet_clean, include_adult="false")
        st = api("/search/tv", query=sheet_clean, first_air_date_year=yr) if yr \
            else api("/search/tv", query=sheet_clean)
        cand = [(r_, "movie") for r_ in sm.get("results", [])[:3]]
        cand += [(r_, "series") for r_ in st.get("results", [])[:3]]
        best, best_score = None, 0.0
        for r_, mtype in cand:
            name = r_.get("title") or r_.get("name")
            score = similar(sheet_clean, name)
            if yr:
                ry = r_.get("release_date") or r_.get("first_air_date") or ""
                try:
                    if int(ry[:4]) == yr:
                        score += 0.3
                except Exception:
                    pass
            if score > best_score:
                best, best_score = (r_, mtype), score
        if best and best_score >= 0.55:
            r_, mtype = best
            rec = build_rec(r_, mtype, imdb)
            return rec, "search", f"search match score={best_score:.2f}"
    except Exception:
        pass
    return None, "failed", ""


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["To See"]
    start_col = 6
    for j, h in enumerate(NEW_COLS):
        ws.cell(row=1, column=start_col + j, value=h)

    manifest = []
    problems = []
    counts = {"imdb": 0, "search": 0, "imdb_mismatch": 0, "failed": 0}

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        title = (row[0] or "").strip()
        year = row[1]
        imdb = imdb_id_from(row[2])
        rec, status, note = resolve(title, year, imdb)
        counts[status] = counts.get(status, 0) + 1
        if rec is None:
            problems.append({"row": i, "sheet_title": title, "year": str(year), "imdb": imdb or "", "status": status})
            vals = ["", "", "", "", imdb or "", "", "", "failed", note]
        else:
            if rec["type"] == "series":
                try:
                    d = api(f"/tv/{rec['tmdb_id']}")
                    rec["seasons"] = d.get("number_of_seasons", "")
                    rec["status"] = d.get("status", "")
                    time.sleep(0.25)
                except Exception:
                    rec["seasons"] = ""
            vals = [
                rec["type"], rec["title"], rec["year"], rec["tmdb_id"],
                rec.get("imdb_id") or imdb or "", rec.get("seasons", ""),
                rec["search_core"], status, note,
            ]
            manifest.append({"row": i, "sheet_title": title, **rec, "resolve": status, "note": note})
            if status == "imdb_mismatch":
                problems.append({"row": i, "sheet_title": title, "tmdb_title": rec["title"], "note": note})
        for j, v in enumerate(vals):
            ws.cell(row=i, column=start_col + j, value=v)
        time.sleep(0.25)

    wb.save(XLSX)
    with open(MANIFEST, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=1, ensure_ascii=False)

    total = sum(counts.values())
    print(f"Resolved {total} rows: imdb={counts['imdb']} search={counts['search']} "
          f"mismatch={counts['imdb_mismatch']} failed={counts['failed']}")
    print(f"Manifest: {MANIFEST} ({len(manifest)} records)")
    if problems:
        print("\nNeeds manual review:")
        for p in problems:
            print(f"  row {p['row']}: {p.get('sheet_title')} -> {p.get('tmdb_title', p.get('status'))} ({p.get('note', '')})")


if __name__ == "__main__":
    main()
